import base64
import copy
import io
import json
import logging
import math
import os
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.table import _Row
from dotenv import find_dotenv, load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from openai import OpenAI
from pydantic import BaseModel

# backend → 상위 폴더 순으로 가장 가까운 .env 로드 (현재: day5/.env)
load_dotenv(find_dotenv())
log = logging.getLogger("uvicorn.error")

BACKEND_DIR = Path(__file__).resolve().parent
FRONTEND_DIR = BACKEND_DIR.parent / "frontend"  # project/frontend
DATA_DIR = BACKEND_DIR.parents[1] / "data"  # module_1/data
TEMPLATE_PATH = DATA_DIR / "report_template.docx"
SAMPLE_FILES = ["img_1.png", "img_2.png", "img_3.png"]

MODEL = "gpt-5.6-luna"
DEFAULT_CF_DB = -50.0
STANDARD = "FCC Part 15 Subpart B Class B (3 m)"
TEST_NAME = "방사성 방출(RE) 측정"  # PRD 미정의 → 템플릿 부제 기반 잠정값 (spec Q2)
REMARK_FIXED = "※ 측정 검출기 RMS(Avg) · Limit 기준 QP — 교육용 참고 판정"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# PRD §3.1 기본 Limit · (freq_min_mhz, freq_max_mhz, limit_dbuv_m) · freq_min ≤ f < freq_max
LIMITS = [
    (30, 88, 40.0),
    (88, 216, 43.5),
    (216, 960, 46.0),
    (960, 6000, 54.0),
]

app = FastAPI(title="QUEST RE API")

# FE 개발 서버(Vite) 오리진 허용
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def api_key_loaded() -> bool:
    return bool(os.getenv("OPENAI_API_KEY", "").strip())


# ---------- 비전 파싱 (PRD F2) ----------

NUM_FIELDS = ("center_freq_ghz", "marker_freq_ghz", "marker_power_uw")
OCR_SCHEMA = {
    "type": "object",
    "properties": {**{k: {"type": "number"} for k in NUM_FIELDS}, "timestamp": {"type": "string"}},
    "required": [*NUM_FIELDS, "timestamp"],
    "additionalProperties": False,
}
OCR_PROMPT = """스펙트럼 아날라이저 화면 캡처에서 아래 4개 값만 추출.
- center_freq_ghz : 하단 좌측 'Center' 주파수 (GHz)
- marker_freq_ghz : 우측 상단 'Mkr1' 주파수 (GHz · 자릿수 공백 제거, 예 2.123 456 → 2.123456)
- marker_power_uw : 'Mkr1' 바로 아래 전력값 (µW)
- timestamp : 상단 날짜·시각 문자열 원문 그대로 (예 01:23:45 PM Jan 01, 2025)
숫자 필드는 단위·설명 없이 숫자만."""


def parse_image(client: OpenAI, name: str, data: bytes) -> dict:
    """스펙트럼 캡처 1장 → Mkr1 주파수·전력 정형 JSON (실패 시 1회 재질의 → OCR_FAIL)"""
    image_url = "data:image/png;base64," + base64.b64encode(data).decode()
    for attempt in (1, 2):
        try:
            res = client.responses.create(
                model=MODEL,
                input=[{"role": "user", "content": [
                    {"type": "input_text", "text": OCR_PROMPT},
                    {"type": "input_image", "image_url": image_url, "detail": "high"},
                ]}],
                # temperature 는 gpt-5.6-luna 미지원(400) → strict JSON 스키마로 출력 고정
                text={"format": {"type": "json_schema", "name": "mkr1_reading", "schema": OCR_SCHEMA, "strict": True}},
            )
            values = json.loads(res.output_text)
            if all(isinstance(values.get(k), (int, float)) and values[k] > 0 for k in NUM_FIELDS):
                return {"file": name, **{k: values[k] for k in OCR_SCHEMA["required"]}, "status": "OK"}
            log.warning("OCR 값 검증 실패 (%s, %d회차): %s", name, attempt, values)
        except Exception as e:  # JSON 파싱 실패 · API 오류
            log.warning("OCR 실패 (%s, %d회차): %s", name, attempt, e)
    return {"file": name, **dict.fromkeys(OCR_SCHEMA["required"]), "status": "OCR_FAIL"}


# ---------- RF 변환 · 판정 (PRD F3 · F4) ----------

def uw_to_dbm(uw: float) -> float:
    return 10 * math.log10(uw / 1000)


def dbm_to_dbuv(dbm: float) -> float:
    return dbm + 107  # 50 Ω 기준


def find_limit(freq_mhz: float) -> float | None:
    return next((limit for lo, hi, limit in LIMITS if lo <= freq_mhz < hi), None)


def summarize(verdicts: list[str]) -> dict:
    if "FAIL" in verdicts:
        overall = "FAIL"
    elif "N/A" in verdicts or not verdicts:
        overall = "INCOMPLETE"
    else:
        overall = "PASS"
    return {"total": len(verdicts), "pass": verdicts.count("PASS"), "fail": verdicts.count("FAIL"), "verdict": overall}


def judge(ocr_rows: list[dict], cf_db: float) -> dict:
    rows = []
    for r in ocr_rows:
        valid = all((r.get(k) or 0) > 0 for k in ("marker_freq_ghz", "marker_power_uw"))
        if r.get("status") != "OK" or not valid:  # OCR 실패 행은 판정 제외 (spec Q5)
            continue
        freq_mhz = round(r["marker_freq_ghz"] * 1000, 6)  # 판정 기준 = Mkr1 주파수 (spec Q4)
        dbm = uw_to_dbm(r["marker_power_uw"])
        dbuv_m = dbm_to_dbuv(dbm) + cf_db
        limit = find_limit(freq_mhz)
        margin = None if limit is None else limit - dbuv_m
        rows.append({
            "no": len(rows) + 1,
            "file": r["file"],
            "timestamp": r.get("timestamp"),
            "freq_mhz": round(freq_mhz, 3),
            "power_uw": round(r["marker_power_uw"], 2),
            "dbm": round(dbm, 2),
            "dbuv_m": round(dbuv_m, 2),
            "limit": limit,
            "margin": None if margin is None else round(margin, 2),
            "verdict": "N/A" if margin is None else "PASS" if margin >= 0 else "FAIL",
        })
    return {"rows": rows, "summary": summarize([row["verdict"] for row in rows]), "cf_db": cf_db}


# ---------- 보고서 렌더링 (PRD F5) ----------

class OcrRow(BaseModel):
    file: str
    center_freq_ghz: float | None = None
    marker_freq_ghz: float | None = None
    marker_power_uw: float | None = None
    timestamp: str | None = None
    status: str


class JudgeRequest(BaseModel):
    ocr: list[OcrRow]
    cf_db: float = DEFAULT_CF_DB


class JudgeRow(BaseModel):
    no: int
    file: str
    timestamp: str | None = None
    freq_mhz: float
    power_uw: float
    dbm: float
    dbuv_m: float
    limit: float | None
    margin: float | None
    verdict: str


class ReportRequest(BaseModel):
    rows: list[JudgeRow]
    cf_db: float
    tester: str = ""
    sample_name: str = ""
    reviewer: str = ""


PLACEHOLDER = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


def fmt(value: float | None, digits: int, sign: bool = False) -> str:
    return "—" if value is None else f"{value:{'+' if sign else ''}.{digits}f}"


def all_paragraphs(container):
    """본문 + 표 셀 문단 (중첩 표 포함)"""
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from all_paragraphs(cell)


def fill(paragraphs, values: dict) -> None:
    for p in paragraphs:
        for run in p.runs:  # run 단위 치환 → 맑은 고딕 등 서식 유지
            if "{{" in run.text:
                run.text = PLACEHOLDER.sub(lambda m: str(values.get(m.group(1), "")), run.text)


def measured_at(rows: list[JudgeRow]) -> datetime:
    """시험일 = 첫 판독 가능 OCR timestamp (spec Q3) · 없으면 생성일"""
    for r in rows:
        try:
            return datetime.strptime(r.timestamp or "", "%I:%M:%S %p %b %d, %Y")
        except ValueError:
            continue
    return datetime.now()


def render_report(req: ReportRequest) -> Document:
    doc = Document(TEMPLATE_PATH)

    # 결과 표 : {{r.*}} 템플릿 행 → 판정 행 수만큼 복제 후 원본 제거
    for table in doc.tables:
        tpl = next((row for row in table.rows if "{{r.no}}" in row.cells[0].text), None)
        if tpl is None:
            continue
        for r in req.rows:
            new_tr = copy.deepcopy(tpl._tr)
            tpl._tr.addprevious(new_tr)
            fill((p for cell in _Row(new_tr, table).cells for p in cell.paragraphs), {
                "r.no": r.no,
                "r.freq_mhz": fmt(r.freq_mhz, 3),
                "r.power_uw": fmt(r.power_uw, 2),
                "r.dbm": fmt(r.dbm, 2),
                "r.dbuv_m": fmt(r.dbuv_m, 2),
                "r.limit": fmt(r.limit, 1),
                "r.margin": fmt(r.margin, 2, sign=True),
                "r.verdict": r.verdict,
            })
        tpl._tr.getparent().remove(tpl._tr)

    dt = measured_at(req.rows)
    summary = summarize([r.verdict for r in req.rows])
    remarks = [f"{r.freq_mhz:.3f} MHz Limit {abs(r.margin):.2f} dB 초과" for r in req.rows if r.verdict == "FAIL"]
    fill(all_paragraphs(doc), {
        "doc_no": f"RE-{dt.year}-0001",
        "test_date": dt.strftime("%Y-%m-%d"),
        "test_name": TEST_NAME,
        "tester": req.tester,
        "sample_name": req.sample_name,
        "reviewer": req.reviewer,
        "standard": STANDARD,
        "cf_db": f"{req.cf_db:.1f}",
        "total": summary["total"],
        "pass_cnt": summary["pass"],
        "fail_cnt": summary["fail"],
        "verdict": summary["verdict"],
        "remarks": "\n".join([*remarks, REMARK_FIXED]),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    })
    return doc


# ---------- API ----------

@app.get("/api/health")
def health():
    return {"status": "ok", "api_key_loaded": api_key_loaded()}


@app.get("/api/samples")
def samples():
    """module_1/data 기본 샘플 목록"""
    return {"files": [n for n in SAMPLE_FILES if (DATA_DIR / n).exists()]}


@app.post("/api/analyze")
def analyze(
    files: list[UploadFile] | None = File(None),
    use_samples: bool = Form(False),
    cf_db: float = Form(DEFAULT_CF_DB),
):
    """이미지(업로드 · 기본 샘플) → 비전 파싱 → RF 변환 · 판정"""
    if not api_key_loaded():
        raise HTTPException(400, "OPENAI_API_KEY 미설정 — .env 에 Key 입력 후 서버 재시작 필요")
    images = [(n, (DATA_DIR / n).read_bytes()) for n in SAMPLE_FILES] if use_samples else []
    images += [(f.filename, f.file.read()) for f in files or []]
    if not images:
        raise HTTPException(400, "분석할 이미지 없음 — PNG 업로드 또는 기본 샘플 로드 필요")
    client = OpenAI(timeout=90)
    with ThreadPoolExecutor(max_workers=4) as pool:
        ocr = list(pool.map(lambda item: parse_image(client, *item), images))
    return {"ocr": ocr, **judge(ocr, cf_db)}


@app.post("/api/judge")
def rejudge(req: JudgeRequest):
    """OCR 결과 재사용 재판정 (CF 변경 시 · OCR 재호출 없음)"""
    return judge([r.model_dump() for r in req.ocr], req.cf_db)


@app.post("/api/report")
def report(req: ReportRequest):
    """report_template.docx 인메모리 렌더링 → re_report.docx 다운로드"""
    buf = io.BytesIO()
    render_report(req).save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type=DOCX_MIME, headers={"Content-Disposition": 'attachment; filename="re_report.docx"'})


# 대시보드 : frontend/index.html 을 루트(/)로 서빙 (API 라우트 등록 이후 마운트)
app.mount("/", StaticFiles(directory=FRONTEND_DIR, html=True), name="frontend")
